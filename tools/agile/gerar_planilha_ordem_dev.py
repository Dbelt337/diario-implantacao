"""Gera a planilha de ordem de desenvolvimento das works do programa: leads, oportunidades, cotacao e pedido.

Uso (na raiz do repo, org btp-prod autenticada):
    python tools/agile/gerar_planilha_ordem_dev.py
    python tools/agile/gerar_planilha_ordem_dev.py --json caminho/consulta.json   (usa um retorno salvo do sf data query)

Saida: docs/2026-09-23-ordem-de-desenvolvimento-works.xlsx
Abas: "Ordem de inicio" (fundacoes, leads, oportunidades, cotacao, pedido), "Leads" (todas as works de leads com narrativa)
e "Todas as works" (107 works classificadas por etapa da jornada).
"""
import json
import os
import re
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ORG = 'btp-prod'
SAIDA = 'docs/2026-09-23-ordem-de-desenvolvimento-works.xlsx'

# Etapa da jornada de cada work e ordem sugerida de inicio dentro da etapa.
# Etapas: 0 Fundacoes minimas, 1 Leads, 2 Oportunidades, 3 Cotacao e contrato, 4 Pedido e ativacao,
# 5 Pos-venda e MACD, 6 Catalogo (paralelo, pre-requisito da cotacao), 7 Tecnicas e Brasil Tecpar (paralelo).
ETAPAS = {
    0: '0 Fundacoes minimas para comecar', 1: '1 Leads', 2: '2 Oportunidades', 3: '3 Cotacao e contrato',
    4: '4 Pedido e ativacao', 5: '5 Pos-venda e MACD', 6: '6 Catalogo (paralelo)', 7: '7 Tecnicas e Brasil Tecpar (paralelo)',
}
ORDEM = [
    # etapa 0: o que precisa existir para a primeira work de leads entrar em desenvolvimento
    (0, 'W-000090', 'Perfis e papeis ja existem na org; a work entrega permission sets, FLS e filas da jornada. Pode andar junto com as primeiras de leads.'),
    (0, 'W-000132', 'Esteira de deploy da SysMap: repositorio e pipeline antes do primeiro merge.'),
    (0, 'W-000133', 'Repositorio GitLab de Brasil Tecpar e esteira basica; sucede a W-000025.'),
    (0, 'W-000087', 'Named Credentials e Integration Procedures base; consulta ao Customer Core e Econodata dependem daqui.'),
    (0, 'W-000140', 'Sandboxes apontando para a homologacao do MuleSoft; sem isso os testes de leads B2B ficam no mock.'),
    # etapa 1: leads
    (1, 'W-000064', 'Empresa do Grupo, Regional e Canal vindos do usuario (SSO/SCIM); base de todos os fluxos.'),
    (1, 'W-000057', 'Captura, triagem e roteamento Omni-Channel do lead B2C; fonte unica de Canal de Entrada.'),
    (1, 'W-000091', 'Filas regionais, skills e presence do Omni-Channel usados pela W-000057.'),
    (1, 'W-000071', 'Ingestao: Web-to-Lead, QR Code, carga massiva restrita por permission set.'),
    (1, 'W-000096', 'Leads B2B: inatividade, escalonamento e enriquecimento via Econodata (job diario no MuleSoft).'),
    (1, 'W-000161', 'Motivos e submotivos de perda (Global Value Set compartilhado Lead e Oportunidade, B2B e B2C).'),
    (1, 'W-000072', 'Busca global federada (Salesforce + Customer Core) usada na qualificacao.'),
    (1, 'W-000073', 'Categorizacao do lead: conversao em Oportunidade ou descarte como perdido.'),
    (1, 'W-000117', 'Modelo de contas (Pessoa Fisica, Pessoa Juridica, Billing) que a conversao cria; Person Account nao habilitado.'),
    (1, 'W-000120', 'Rastreio comercial e identificacao de parceiros (Finder, Integrador) no lead e na oportunidade.'),
    (1, 'W-000083', 'Guarda de dados e consentimento (LGPD) para leads perdidos e abordagens futuras.'),
    # etapa 2: oportunidades (inclui qualificacao e viabilidade, que rodam na oportunidade)
    (2, 'W-000062', 'Tipo de Negociacao e segmentacao B2C/B2S por ticket; campo obrigatorio antes do carrinho.'),
    (2, 'W-000060', 'Pre-viabilidade: debito no endereco antes da consulta tecnica.'),
    (2, 'W-000074', 'Viabilidade tecnica e credito B2C em paralelo (Integration Procedure existente refatorada).'),
    (2, 'W-000058', 'Trilha alternativa: oferta movel quando inviavel.'),
    (2, 'W-000059', 'Regra de risco: contratos sem 1a parcela (Decision Table do BRE).'),
    (2, 'W-000097', 'B2B: enderecamento geocodificado, viabilidade expressa por site (TMF645 via Mule) e roteamento.'),
    (2, 'W-000119', 'B2B: importacao de sites multi-ponto por planilha.'),
    (2, 'W-000066', 'Regua de lembretes e expiracao da oportunidade por SLA (Marketing Cloud; depende da W-000092).'),
    (2, 'W-000160', 'B2B: restricao de edicao pelo BKO e bypass de arquitetura.'),
    (2, 'W-000115', 'B2C: tipos de operacao sem faturamento (troca, demonstracao, cortesia).'),
    (2, 'W-000162', 'B2B: configuracao e precificacao de cortesia.'),
    (2, 'W-000163', 'B2B: aprovacao mandatoria e trava comercial para cortesias (Flow Approval Orchestration).'),
    (2, 'W-000165', 'B2B: permuta de servicos (Swap/SUAP).'),
    (2, 'W-000169', 'B2C: venda cortesia com aprovacao previa.'),
    (2, 'W-000170', 'B2C: venda swap com preco flexivel.'),
    (2, 'W-000123', 'B2B: visao 360 de contratos e ativos do grupo economico (consulta na oportunidade).'),
    (2, 'W-000061', 'Dashboard gerencial e arvore de perdas (usa os motivos da W-000161).'),
    # etapa 3: cotacao e contrato
    (3, 'W-000084', 'Botao Criar Cotacao na Oportunidade (Industries CPQ Create Cart).'),
    (3, 'W-000086', 'Esqueleto do OmniScript Nova Venda B2C.'),
    (3, 'W-000069', 'Selecao guiada, combos e filtro por IBGE (depende do catalogo, etapa 6).'),
    (3, 'W-000107', 'Motor de regras de precificacao de projetos especiais (BRE).'),
    (3, 'W-000108', 'Regra de desconto em combos e promocoes.'),
    (3, 'W-000121', 'Modalidade de pagamento unico (CAPEX).'),
    (3, 'W-000098', 'B2B: cotacao multi-site, alcadas de desconto e proposta.'),
    (3, 'W-000144', 'B2B: condicoes especiais de faturamento.'),
    (3, 'W-000068', 'B2C: taxa de ativacao e alcadas de desconto.'),
    (3, 'W-000081', 'B2C: dados de faturamento e parametros de cobranca.'),
    (3, 'W-000089', 'Infraestrutura de Document Generation.'),
    (3, 'W-000106', 'Templates juridicos unificados.'),
    (3, 'W-000082', 'Minuta do contrato B2C gerada da cotacao.'),
    (3, 'W-000122', 'B2B: signatario legal, procuracao e contato de NPS no contrato.'),
    (3, 'W-000065', 'Resumo da venda e modalidade de aceite (sem motor de assinatura licenciado: aceite por evidencia).'),
    (3, 'W-000105', 'Canal de eventos de contrato Salesforce para MuleSoft.'),
    (3, 'W-000145', 'B2B: cadencia de notificacoes de assinatura (Marketing Cloud).'),
    (3, 'W-000092', 'Fundacao Marketing Cloud (jornadas base); WhatsApp depende de add-on a confirmar.'),
    (3, 'W-000075', 'Recepcao inbound de vendas digitais (Touchless Order API).'),
    # etapa 4: pedido e ativacao (bloqueada em producao ate a licenca de OM, W-000134)
    (4, 'W-000118', 'Checkout da cotacao e submissao do pedido ao Order Management (gatilho B2C e B2B).'),
    (4, 'W-000088', 'Decomposicao e plano de orquestracao B2C no Industries OM.'),
    (4, 'W-000104', 'Algoritmo da Service Tag no pedido.'),
    (4, 'W-000101', 'B2B: auditoria pelo BKO, credito e handoff fiscal.'),
    (4, 'W-000067', 'B2C: debito interno, ticket Zendesk e SLA.'),
    (4, 'W-000076', 'B2C: emissao da cobranca, comprovante e Mesa de Credito.'),
    (4, 'W-000077', 'B2C: controle de pagamento, isencao e autoagendamento.'),
    (4, 'W-000078', 'B2C: list view de acompanhamento da baixa bancaria.'),
    (4, 'W-000093', 'B2C: orquestracao de ativacao e provisionamento no Customer Core.'),
    (4, 'W-000094', 'B2C: Work Order de instalacao (OM + Field Service).'),
    (4, 'W-000063', 'Field Service: reagendamento e cancelamento por insucesso.'),
    (4, 'W-000079', 'Relatorio de instalacoes nao realizadas e notificacao.'),
    (4, 'W-000080', 'Ativacao final no Customer Core e fechamento como Ganho.'),
    (4, 'W-000110', 'Integracao com fornecedores de SVA (item callout do OM).'),
    (4, 'W-000130', 'Sincronizacao SAP para Salesforce do equipamento instalado.'),
    (4, 'W-000131', 'API inbound TMF641 de parceiros de delivery.'),
    (4, 'W-000164', 'B2B: decomposicao e handoff de cortesia (faturamento nulo).'),
    (4, 'W-000168', 'B2C: relatorio de auditoria de cortesia e swap.'),
    # etapa 5: pos-venda e MACD
    (5, 'W-000095', ''), (5, 'W-000099', ''), (5, 'W-000100', ''), (5, 'W-000109', ''), (5, 'W-000116', ''), (5, 'W-000124', ''),
    (5, 'W-000125', ''), (5, 'W-000126', ''), (5, 'W-000128', ''), (5, 'W-000129', ''), (5, 'W-000142', ''), (5, 'W-000143', ''),
    (5, 'W-000166', ''), (5, 'W-000167', ''), (5, 'W-000171', ''), (5, 'W-000172', ''), (5, 'W-000173', ''),
    # etapa 6: catalogo (paralelo; pre-requisito da etapa 3)
    (6, 'W-000051', ''), (6, 'W-000052', ''), (6, 'W-000053', ''), (6, 'W-000054', ''), (6, 'W-000055', ''), (6, 'W-000056', ''),
    (6, 'W-000070', ''), (6, 'W-000085', ''), (6, 'W-000102', ''), (6, 'W-000103', ''), (6, 'W-000111', ''), (6, 'W-000112', ''),
    (6, 'W-000113', ''), (6, 'W-000114', ''), (6, 'W-000127', ''),
    # etapa 7: tecnicas e Brasil Tecpar (paralelo)
    (7, 'W-000134', 'Licenca de OM: bloqueia a etapa 4 em producao.'), (7, 'W-000135', ''), (7, 'W-000136', ''), (7, 'W-000137', ''),
    (7, 'W-000141', ''),
]
LEADS = [n for e, n, _ in ORDEM if e == 1]
INICIO = [0, 1, 2, 3, 4]

FONTE = 'Arial'
CAB_FILL = PatternFill('solid', fgColor='1F3864')
ETAPA_FILL = {0: 'E7E6E6', 1: 'DDEBF7', 2: 'E2EFDA', 3: 'FFF2CC', 4: 'FCE4D6', 5: 'EDEDED', 6: 'EDEDED', 7: 'EDEDED'}
BORDA = Border(*(Side(style='thin', color='BFBFBF'),) * 4)


def carregar(json_path=None):
    if json_path:
        return json.load(open(json_path, encoding='utf8'))['result']['records']
    soql = ("SELECT Id, Name, agf__Subject__c, agf__Details__c, agf__Priority__c, agf__Status__c, agf__Epic__r.Name, "
            "agf__Sprint__r.Name, agf__Product_Owner__r.Name, agf__QA_Engineer__r.Name, agf__Assignee__r.Name "
            "FROM agf__ADM_Work__c WHERE agf__Subject__c LIKE 'US %' ORDER BY Name")
    with tempfile.NamedTemporaryFile('w', suffix='.soql', delete=False, encoding='utf8') as f:
        f.write(soql)
        p = f.name
    out = subprocess.run(['sf', 'data', 'query', '--file', p, '--target-org', ORG, '--json'],
                         capture_output=True, text=True, encoding='utf8', shell=(os.name == 'nt'))
    os.unlink(p)
    return json.loads(out.stdout)['result']['records']


def secao(texto, nomes):
    """Devolve o texto da primeira secao cujo titulo (linha em maiusculas) esteja em nomes."""
    linhas = (texto or '').replace('\r\n', '\n').split('\n')
    for i, l in enumerate(linhas):
        if l.strip().rstrip(':') in nomes:
            corpo = []
            for m in linhas[i + 1:]:
                s = m.strip()
                if s and s.isupper() and len(s) < 60 and ' ' in s:
                    break
                if s:
                    corpo.append(s)
            return ' '.join(corpo)
    return ''


def narrativa(texto):
    t = secao(texto, {'NARRATIVA', '1. NARRATIVA DE NEGÓCIO', '1. NARRATIVA DE NEGOCIO', 'NARRATIVA DE NEGÓCIO'})
    if not t:
        m = re.search(r'(?s)Como .*?(?:\n\n|$)', texto or '')
        t = m.group(0).strip() if m else ''
    return t[:700]


def dependencias(texto, nome):
    dep = secao(texto, {'DEPENDENCIAS E LICENCAS', 'DEPENDÊNCIAS E LICENÇAS', 'DEPENDENCIAS', 'DEPENDÊNCIAS'})
    if not dep:
        m = re.search(r'Depend[eê]ncias?:([^\n]*)', texto or '')
        dep = m.group(1) if m else ''
    refs = sorted(set(re.findall(r'W-\d{6}', dep)) - {nome})
    if not refs:
        refs = sorted(set(re.findall(r'W-\d{6}', texto or '')) - {nome})
        return ', '.join(refs), 'texto todo'
    return ', '.join(refs), 'secao de dependencias'


def rel(r, campo):
    return (r.get(campo) or {}).get('Name') or ''


def cabecalho(ws, cols, larguras):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(name=FONTE, bold=True, color='FFFFFF', size=10)
        cell.fill = CAB_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = larguras[i - 1]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}1'
    ws.row_dimensions[1].height = 30


def linha(ws, valores, fill=None):
    ws.append(valores)
    r = ws.max_row
    for i in range(1, len(valores) + 1):
        c = ws.cell(row=r, column=i)
        c.font = Font(name=FONTE, size=10)
        c.alignment = Alignment(vertical='top', wrap_text=True)
        c.border = BORDA
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)


def gerar(records, saida):
    by = {r['Name']: r for r in records}
    etapa_de = {n: e for e, n, _ in ORDEM}
    obs_de = {n: o for e, n, o in ORDEM}
    faltam = sorted(set(by) - set(etapa_de))
    if faltam:
        print('works sem etapa (entram como nao classificadas):', faltam)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Ordem de inicio'
    cols = ['Ordem', 'Etapa', 'Work', 'Titulo', 'Prioridade', 'Status', 'Sprint atual', 'Epico', 'Depende de',
            'Origem da dependencia', 'Product Owner', 'Observacao para o inicio']
    cabecalho(ws, cols, [7, 30, 11, 55, 10, 10, 26, 32, 28, 16, 16, 60])
    ordem = 0
    for e, n, o in ORDEM:
        if e not in INICIO or n not in by:
            continue
        r = by[n]
        ordem += 1
        deps, origem = dependencias(r['agf__Details__c'], n)
        linha(ws, [ordem, ETAPAS[e], n, r['agf__Subject__c'], r['agf__Priority__c'] or '', r['agf__Status__c'] or '',
                   rel(r, 'agf__Sprint__r'), rel(r, 'agf__Epic__r'), deps, origem, rel(r, 'agf__Product_Owner__r'), o], ETAPA_FILL[e])
    ws.append([])
    ws.append(['Legenda: a ordem e uma sugestao de sequencia de desenvolvimento; a coluna Sprint atual mostra a alocacao de hoje no Agile '
               'para comparar. Depende de lista as works citadas na secao de dependencias do texto (ou no texto todo, quando a secao nao existe). '
               'Etapa 4 so vai a producao apos a licenca de Order Management (W-000134). Catalogo (etapa 6) e works tecnicas (etapa 7) '
               'ficam na aba Todas as works.'])
    ws.cell(row=ws.max_row, column=1).font = Font(name=FONTE, size=9, italic=True)
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=12)
    ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[ws.max_row].height = 45

    ws2 = wb.create_sheet('Leads')
    cols = ['Ordem', 'Work', 'Titulo', 'Prioridade', 'Status', 'Sprint atual', 'Epico', 'Product Owner', 'QA Engineer',
            'Narrativa', 'Depende de', 'Dependencias e licencas (texto da work)', 'Observacao para o inicio']
    cabecalho(ws2, cols, [7, 11, 50, 10, 10, 26, 30, 16, 16, 70, 24, 60, 50])
    for i, n in enumerate(LEADS, 1):
        r = by.get(n)
        if not r:
            continue
        deps, _ = dependencias(r['agf__Details__c'], n)
        dep_txt = secao(r['agf__Details__c'], {'DEPENDENCIAS E LICENCAS', 'DEPENDÊNCIAS E LICENÇAS', 'DEPENDENCIAS', 'DEPENDÊNCIAS'})[:600]
        linha(ws2, [i, n, r['agf__Subject__c'], r['agf__Priority__c'] or '', r['agf__Status__c'] or '', rel(r, 'agf__Sprint__r'),
                    rel(r, 'agf__Epic__r'), rel(r, 'agf__Product_Owner__r'), rel(r, 'agf__QA_Engineer__r'),
                    narrativa(r['agf__Details__c']), deps, dep_txt, obs_de.get(n, '')], ETAPA_FILL[1])

    ws3 = wb.create_sheet('Todas as works')
    cols = ['Etapa', 'Ordem na etapa', 'Work', 'Titulo', 'Prioridade', 'Status', 'Sprint atual', 'Epico', 'Product Owner',
            'QA Engineer', 'Assignee', 'Depende de', 'Observacao']
    cabecalho(ws3, cols, [30, 9, 11, 55, 10, 10, 26, 32, 16, 16, 16, 28, 55])
    pos = {}
    for e, n, o in ORDEM:
        pos[n] = (e, len([1 for x in ORDEM[:ORDEM.index((e, n, o))] if x[0] == e]) + 1)
    for n in sorted(by, key=lambda x: (pos.get(x, (9, 0)), x)):
        r = by[n]
        e, k = pos.get(n, (9, 0))
        deps, _ = dependencias(r['agf__Details__c'], n)
        linha(ws3, [ETAPAS.get(e, '9 Nao classificada'), k or '', n, r['agf__Subject__c'], r['agf__Priority__c'] or '',
                    r['agf__Status__c'] or '', rel(r, 'agf__Sprint__r'), rel(r, 'agf__Epic__r'), rel(r, 'agf__Product_Owner__r'),
                    rel(r, 'agf__QA_Engineer__r'), rel(r, 'agf__Assignee__r'), deps, obs_de.get(n, '')], ETAPA_FILL.get(e, 'FFFFFF'))
    wb.save(saida)
    return ordem, len(LEADS), len(by)


if __name__ == '__main__':
    jp = sys.argv[sys.argv.index('--json') + 1] if '--json' in sys.argv else None
    recs = carregar(jp)
    o, l, t = gerar(recs, SAIDA)
    print(f'gerado: {SAIDA} | ordem de inicio: {o} works | leads: {l} | total: {t}')
